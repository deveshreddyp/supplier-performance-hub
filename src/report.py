"""Generate formatted Excel supplier reports."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from src.analyze import analyze
from src.db import frame
ROOT=Path(__file__).resolve().parents[1]
def generate_report(path=None):
    path=Path(path or ROOT/"reports/supplier_report.xlsx"); path.parent.mkdir(parents=True,exist_ok=True); a=analyze()
    issues=frame("SELECT severity,table_name,record_key,issue,created_at FROM validation_report ORDER BY id DESC")
    wb=Workbook(); wb.remove(wb.active)
    for name,df in {"Scorecards":a["scorecards"],"Breaches":a["breaches"],"Renewals":a["renewals"],"Validation Issues":issues}.items():
        ws=wb.create_sheet(name); ws.append(list(df.columns) if not df.empty else ["No records"])
        for row in df.itertuples(index=False,name=None):
            ws.append([v.isoformat() if hasattr(v,"isoformat") else (None if pd.isna(v) else v) for v in row])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D"); cell.alignment=Alignment(wrap_text=True)
        for col in ws.columns: ws.column_dimensions[get_column_letter(col[0].column)].width=min(42,max(12,max(len(str(c.value or "")) for c in col)+2))
    wb.save(path); return path
if __name__=="__main__": print(generate_report())
